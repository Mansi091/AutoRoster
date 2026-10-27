"""Generate bus route stops using the centralized route_planner and settings.json.

Enforces 32-seat bus capacity and optimal stop sequencing terminating at the factory.
"""
from pathlib import Path
import sys
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

import route_planner as rp
import plan_routes as pr

ROOT = Path(__file__).resolve().parent

def main():
    settings = pr.load_settings(ROOT / "settings.json")
    settings["bus_capacity"] = 32
    roster_path = ROOT / settings["roster_path"]
    stops_path = ROOT / settings["stops_path"]

    fac = settings["facility"]
    p = rp.RoutingParams(
        bus_capacity=32,
        max_ride_minutes=float(settings["max_ride_minutes"]),
        bus_fixed_cost_minutes=float(settings["bus_fixed_cost_minutes"]),
        boarding_seconds_per_passenger=float(settings["boarding_seconds_per_passenger"]),
        min_boarding_minutes=float(settings["min_boarding_minutes"]),
    )

    employees, _ = pr.read_roster(roster_path, settings.get("roster_sheet"), settings)
    stops = pr.read_stop_master(stops_path, settings.get("stops_sheet"), settings)
    priority = [s.upper() for s in settings["shift_priority"]]
    rows, _ = pr.assign_employees(employees, priority, stops)

    points = [(fac["latitude"], fac["longitude"])]
    node_of = {}
    for key in sorted({r["pickup_key"] for r in rows if r["pickup_key"] in stops}):
        coord = (stops[key]["lat"], stops[key]["lon"])
        if coord not in node_of:
            node_of[coord] = len(points)
            points.append(coord)

    matrix = rp.build_travel_matrix(
        points, settings["travel_time_provider"], settings["osrm_url"],
        float(settings["fallback_road_factor"]), float(settings["fallback_speed_kmph"])
    )
    M = matrix.minutes
    FAC = 0

    route_overview = []
    detailed_stops = []

    for shift in priority:
        units = pr.build_units(rows, shift, stops, node_of, p)
        if not units:
            continue
        routes = rp.group_units(units, M, FAC, p)

        for ridx, route in enumerate(routes, start=1):
            bus_id = f"Bus {shift}-{ridx:02d}"
            rides, _ = rp.evaluate(route, units, M, FAC)
            seq_nodes = [units[u].node for u in route]
            total_pax = sum(units[u].pax for u in route)
            road_km = sum(matrix.km[seq_nodes[k]][seq_nodes[k + 1] if k + 1 < len(seq_nodes) else FAC]
                          for k in range(len(seq_nodes)))

            first_stop = units[route[0]].stop_name
            last_stop = units[route[-1]].stop_name

            cum_km = 0.0
            cum_pax = 0
            prev_node = None
            max_jump = 0.0

            for pos, u in enumerate(route, 1):
                unit = units[u]
                dist_prev = matrix.km[prev_node][unit.node] if prev_node is not None else 0.0
                cum_km += dist_prev
                cum_pax += unit.pax
                if dist_prev > max_jump:
                    max_jump = dist_prev

                detailed_stops.append({
                    "Bus ID": bus_id,
                    "Shift": f"Shift {shift}",
                    "Stop Seq": pos,
                    "Bus Stop Name": unit.stop_name,
                    "Latitude": round(unit.stop_key[0], 6),
                    "Longitude": round(unit.stop_key[1], 6),
                    "Distance from Prev Stop (km)": round(dist_prev, 1),
                    "Cumulative Route (km)": round(cum_km, 1),
                    "Employees Boarding": unit.pax,
                    "Cumulative Passengers": cum_pax,
                })
                prev_node = unit.node

            route_overview.append({
                "Bus ID": bus_id,
                "Shift": f"Shift {shift}",
                "Target Capacity": 32,
                "Passengers Assigned": total_pax,
                "Seats Free": 32 - total_pax,
                "Total Stops": len(route),
                "Total Route (km)": round(road_km, 1),
                "Max Distance Between Stops (km)": round(max_jump, 1),
                "Start Stop": first_stop,
                "End Stop": last_stop,
            })

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Bus Routes (32 Pax)"
    ws1.views.sheetView[0].showGridLines = True

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    regular_font = Font(name="Calibri", size=11)
    bold_font = Font(name="Calibri", size=11, bold=True)
    thin_border = Border(left=Side(style="thin", color="D9D9D9"),
                         right=Side(style="thin", color="D9D9D9"),
                         top=Side(style="thin", color="D9D9D9"),
                         bottom=Side(style="thin", color="D9D9D9"))

    ov_df = pd.DataFrame(route_overview)
    ws1.append(list(ov_df.columns))
    for cell in ws1[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in ov_df.to_dict(orient="records"):
        ws1.append(list(r.values()))

    ws2 = wb.create_sheet(title="Stop-by-Stop Route Sequence")
    ws2.views.sheetView[0].showGridLines = True
    det_df = pd.DataFrame(detailed_stops)
    ws2.append(list(det_df.columns))
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in det_df.to_dict(orient="records"):
        ws2.append(list(r.values()))

    for ws in (ws1, ws2):
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    out_file = ROOT / "Bus_Stops_Grouped_By_Route.xlsx"
    wb.save(out_file)
    print(f"Generated: {out_file} (Capacity 32, {len(route_overview)} buses)")

if __name__ == "__main__":
    main()
