"""Plan employee bus routes from a roster and a verified pickup-point master.

Usage
-----
    python plan_routes.py                     # paths and rules from settings.json
    python plan_routes.py --roster raw.xlsx --stops master.xlsx --out plan.xlsx

Everything location-, time- or rule-specific comes from settings.json or from
the input files. Changing the raw roster or moving the drop point requires no
code change.
"""

from __future__ import annotations

import argparse
import json
import logging
import math
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import route_planner as rp

ROOT = Path(__file__).resolve().parent
LOGGER = logging.getLogger("plan_routes")


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def load_settings(path: Path) -> dict:
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def norm_key(text) -> str:
    """Canonical key for matching names: lower case, punctuation removed.

    'Main Rd, Area A' and 'main-rd area a' give the same key.
    """
    return re.sub(r"[^a-z0-9]+", " ", str(text).lower()).strip()


def clock_to_minutes(text: str) -> int:
    hh, mm = str(text).strip().split(":")
    return int(hh) * 60 + int(mm)


def minutes_to_clock(value: float) -> str:
    total = int(round(value)) % (24 * 60)
    return f"{total // 60:02d}:{total % 60:02d}"


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, float) and math.isnan(value)) or str(value).strip() == ""


def clean_mobile(value) -> str:
    """Keep the number as text. Excel often stores mobiles as floats, which drops a leading zero."""
    if _is_blank(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return re.sub(r"\s+", "", text)


def parse_shift_tags(value, priority: List[str]) -> Tuple[List[str], List[str]]:
    tokens = [t for t in re.split(r"[\s+,&/]+", str(value).upper()) if t]
    known = [t for t in tokens if t in priority]
    unknown = [t for t in tokens if t not in priority]
    return known, unknown


# ---------------------------------------------------------------------------
# Reading the roster (header row and column names are detected, not fixed)
# ---------------------------------------------------------------------------

def _match_columns(headers: List[str], aliases: Dict[str, List[str]]) -> Dict[str, int]:
    found: Dict[str, int] = {}
    for std, names in aliases.items():
        wanted = {norm_key(n) for n in names}
        for idx, header in enumerate(headers):
            if header in wanted:
                found[std] = idx
                break
    return found


def read_roster(path: Path, sheet, settings: dict) -> Tuple[List[dict], List[dict]]:
    aliases = settings["roster_columns"]
    raw = pd.read_excel(path, sheet_name=sheet if sheet is not None else 0, header=None, dtype=object)

    header_row, columns = None, None
    for r in range(min(40, len(raw))):
        headers = [norm_key(v) if not _is_blank(v) else "" for v in raw.iloc[r].tolist()]
        found = _match_columns(headers, aliases)
        if {"employee_name", "pickup_point", "shift"} <= found.keys():
            header_row, columns = r, found
            break
    if header_row is None:
        raise ValueError(
            f"Could not find a header row with name, pickup and shift columns in {path}. "
            f"Checked aliases in settings.json -> roster_columns."
        )

    employees, issues = [], []
    for excel_row in range(header_row + 1, len(raw)):
        row = raw.iloc[excel_row]
        if all(_is_blank(v) for v in row.tolist()):
            continue

        def cell(name):
            idx = columns.get(name)
            return None if idx is None else row.iloc[idx]

        employees.append({
            "source_row": excel_row + 1,
            "employee_id": "" if _is_blank(cell("employee_id")) else str(cell("employee_id")).strip(),
            "employee_name": "" if _is_blank(cell("employee_name")) else str(cell("employee_name")).strip(),
            "shift_raw": "" if _is_blank(cell("shift")) else str(cell("shift")).strip(),
            "pickup_raw": "" if _is_blank(cell("pickup_point")) else str(cell("pickup_point")).strip(),
            "mobile": clean_mobile(cell("mobile")),
            "area": "" if _is_blank(cell("area")) else str(cell("area")).strip(),
        })

    LOGGER.info("Read %d roster rows from header row %d of %s", len(employees), header_row + 1, path.name)
    return employees, issues


# ---------------------------------------------------------------------------
# Reading the verified pickup-point master
# ---------------------------------------------------------------------------

def read_stop_master(path: Path, sheet, settings: dict) -> Dict[str, dict]:
    aliases = settings["stop_columns"]
    verified = {s.upper() for s in settings["verified_statuses"]}
    df = pd.read_excel(path, sheet_name=sheet if sheet is not None else 0, dtype=object)
    headers = [norm_key(c) for c in df.columns]
    found = _match_columns(headers, aliases)
    for required in ("name", "latitude", "longitude"):
        if required not in found:
            raise ValueError(f"Stop master {path} has no column for '{required}' (see settings.json -> stop_columns).")

    def col(row, key):
        idx = found.get(key)
        return None if idx is None else row.iloc[idx]

    stops: Dict[str, dict] = {}
    ambiguous: set = set()
    for i in range(len(df)):
        row = df.iloc[i]
        name = col(row, "name")
        if _is_blank(name):
            continue
        status = "" if _is_blank(col(row, "status")) else str(col(row, "status")).strip().upper()
        try:
            lat = float(col(row, "latitude"))
            lon = float(col(row, "longitude"))
        except (TypeError, ValueError):
            continue
        if not (-90 <= lat <= 90 and -180 <= lon <= 180) or (lat == 0 and lon == 0):
            continue
        if status not in verified:
            continue  # only verified pins are routed; the rest are reported as unverified

        info = {"name": str(name).strip(), "lat": lat, "lon": lon, "status": status}
        keys = [norm_key(name)]
        alias_text = col(row, "alias")
        if not _is_blank(alias_text):
            keys += [norm_key(a) for a in str(alias_text).split("|") if norm_key(a)]
        for key in keys:
            if key in stops and (abs(stops[key]["lat"] - lat) > 1e-6 or abs(stops[key]["lon"] - lon) > 1e-6):
                ambiguous.add(key)
            stops.setdefault(key, info)

    for key in ambiguous:  # same name, different pins: do not guess
        stops.pop(key, None)
        LOGGER.warning("Pickup name '%s' has conflicting verified coordinates and was excluded.", key)
    return stops


# ---------------------------------------------------------------------------
# Building units and routes
# ---------------------------------------------------------------------------

def assign_employees(employees: List[dict], priority: List[str], stops: Dict[str, dict]):
    """Attach shift, stop and data-quality flags to each employee. Returns (rows, review_items)."""
    rows, review = [], []
    for emp in employees:
        known, unknown = parse_shift_tags(emp["shift_raw"], priority)
        key = norm_key(emp["pickup_raw"])
        rec = dict(emp, pickup_key=key, shift=None, shift_tags=known)
        if unknown:
            review.append(_review("Unknown shift tag", emp, f"Shift '{emp['shift_raw']}' contains unknown tag(s) {unknown}"))
        if not known:
            review.append(_review("No usable shift", emp, f"Shift '{emp['shift_raw']}' does not match any configured shift"))
        else:
            rec["shift"] = min(known, key=priority.index)  # multi-shift: earliest configured shift
        if not emp["employee_id"]:
            review.append(_review("Missing employee ID", emp, "Employee ID is blank in the roster"))
        if emp["mobile"] and not re.fullmatch(r"\d{10}", emp["mobile"]):
            review.append(_review("Mobile number", emp, f"Mobile '{emp['mobile']}' is not 10 digits. Check for a lost leading zero."))
        if not emp["mobile"]:
            review.append(_review("Mobile number", emp, "Mobile number is blank"))
        if key not in stops:
            review.append(_review("Unverified pickup", emp, f"No verified coordinate for pickup '{emp['pickup_raw']}'. Not routed."))
        rows.append(rec)
    return rows, review


def _review(category: str, emp: dict, detail: str, shift: Optional[str] = None) -> dict:
    return {
        "Category": category,
        "Shift": shift or "",
        "Source row": emp.get("source_row", ""),
        "Employee ID": emp.get("employee_id", ""),
        "Employee Name": emp.get("employee_name", ""),
        "Pickup Point": emp.get("pickup_raw", ""),
        "Detail": detail,
    }


def build_units(rows: List[dict], shift: str, stops: Dict[str, dict], node_of: Dict[Tuple[float, float], int],
                p: rp.RoutingParams) -> List[rp.Unit]:
    """One unit per physical pin for this shift. Names sharing a pin are one bus stop."""
    by_pin: Dict[Tuple[float, float], List[dict]] = {}
    for rec in rows:
        if rec["shift"] == shift and rec["pickup_key"] in stops:
            info = stops[rec["pickup_key"]]
            by_pin.setdefault((info["lat"], info["lon"]), []).append(rec)

    units: List[rp.Unit] = []
    for pin in sorted(by_pin):
        members = by_pin[pin]
        names = sorted({stops[r["pickup_key"]]["name"] for r in members})
        label = " / ".join(names)
        for start in range(0, len(members), p.bus_capacity):
            part = members[start:start + p.bus_capacity]
            pax = len(part)
            dwell = max(p.min_boarding_minutes, pax * p.boarding_seconds_per_passenger / 60.0)
            name = label if len(members) <= p.bus_capacity else f"{label} (part {start // p.bus_capacity + 1})"
            units.append(rp.Unit(unit_id=len(units), stop_key=pin, stop_name=name,
                                 node=node_of[pin], pax=pax, dwell_min=dwell, employees=part))
    return units


def shared_pin_review(stops: Dict[str, dict]) -> List[dict]:
    """Different place names pinned to the same coordinates are a likely geocoding error."""
    by_pin: Dict[Tuple[float, float], set] = {}
    for info in stops.values():
        by_pin.setdefault((info["lat"], info["lon"]), set()).add(info["name"])
    out = []
    for pin, names in by_pin.items():
        if len(names) > 1:
            out.append({"Category": "Shared pin: check coordinates", "Shift": "", "Source row": "",
                        "Employee ID": "", "Employee Name": "", "Pickup Point": " / ".join(sorted(names)),
                        "Detail": f"{len(names)} different names have the identical pin {pin[0]:.6f}, {pin[1]:.6f}"})
    return out


# ---------------------------------------------------------------------------
# Main planning
# ---------------------------------------------------------------------------

def plan(args) -> None:
    settings = load_settings(args.settings)
    roster_path = Path(args.roster or settings["roster_path"])
    stops_path = Path(args.stops or settings["stops_path"])
    out_path = Path(args.out or settings["output_path"])
    roster_path = roster_path if roster_path.is_absolute() else ROOT / roster_path
    stops_path = stops_path if stops_path.is_absolute() else ROOT / stops_path
    out_path = out_path if out_path.is_absolute() else ROOT / out_path

    fac = settings["facility"]
    depot = settings.get("depot")
    priority = [s.upper() for s in settings["shift_priority"]]
    shift_times = {k.upper(): clock_to_minutes(v) for k, v in settings["shifts"].items()}
    buffer_min = float(settings["arrival_buffer_minutes"])

    p = rp.RoutingParams(
        bus_capacity=int(settings["bus_capacity"]),
        max_ride_minutes=float(settings["max_ride_minutes"]),
        bus_fixed_cost_minutes=float(settings["bus_fixed_cost_minutes"]),
        boarding_seconds_per_passenger=float(settings["boarding_seconds_per_passenger"]),
        min_boarding_minutes=float(settings["min_boarding_minutes"]),
    )

    employees, _ = read_roster(roster_path, args.roster_sheet or settings.get("roster_sheet"), settings)
    stops = read_stop_master(stops_path, args.stops_sheet or settings.get("stops_sheet"), settings)
    rows, review = assign_employees(employees, priority, stops)
    review.extend(shared_pin_review(stops))

    # One travel matrix for every point we might visit, so each stop is queried once.
    points: List[Tuple[float, float]] = [(fac["latitude"], fac["longitude"])]
    if depot:
        points.append((depot["latitude"], depot["longitude"]))
    node_of: Dict[Tuple[float, float], int] = {}
    for key in sorted({r["pickup_key"] for r in rows if r["pickup_key"] in stops}):
        coord = (stops[key]["lat"], stops[key]["lon"])
        if coord not in node_of:
            node_of[coord] = len(points)
            points.append(coord)
    matrix = rp.build_travel_matrix(points, settings["travel_time_provider"], settings["osrm_url"],
                                    float(settings["fallback_road_factor"]), float(settings["fallback_speed_kmph"]))
    LOGGER.info("Travel matrix: %d points, source = %s", len(points), matrix.source)
    M = matrix.minutes
    FAC = 0
    DEP = 1 if depot else None

    summary_rows, stop_rows, passenger_rows = [], [], []
    for shift in priority:
        if shift not in shift_times:
            review.append({"Category": "Shift not configured", "Shift": shift, "Source row": "", "Employee ID": "",
                           "Employee Name": "", "Pickup Point": "", "Detail": f"No start time for shift {shift} in settings.json"})
            continue
        units = build_units(rows, shift, stops, node_of, p)
        if not units:
            continue
        arrival = shift_times[shift] - buffer_min  # latest arrival at the facility, minutes from midnight

        routes = rp.group_units(units, M, FAC, p)
        LOGGER.info("Shift %s: %d passengers, %d stops, %d buses", shift, sum(u.pax for u in units), len(units), len(routes))

        for ridx, route in enumerate(routes, start=1):
            bus_id = f"{shift}-{ridx:02d}"
            rides, pax_min = rp.evaluate(route, units, M, FAC)
            seq_nodes = [units[u].node for u in route]
            dry_min = M[DEP][seq_nodes[0]] if DEP is not None else 0.0
            dry_km = matrix.km[DEP][seq_nodes[0]] if DEP is not None else 0.0
            route_min = rides[0] if rides else 0.0
            depart = arrival - (route_min + dry_min)
            road_km = dry_km + sum(matrix.km[seq_nodes[k]][seq_nodes[k + 1] if k + 1 < len(seq_nodes) else FAC]
                                   for k in range(len(seq_nodes)))
            pax_total = sum(units[u].pax for u in route)
            over = []

            for pos, u in enumerate(route):
                unit = units[u]
                ride = rides[pos]
                reach = arrival - ride
                flag = "" if ride <= p.max_ride_minutes + 1e-9 else f"Ride {ride:.0f} min exceeds {p.max_ride_minutes:.0f}"
                stop_rows.append({
                    "Bus ID": bus_id, "Shift": shift, "Stop Seq": pos + 1, "Stop Name": unit.stop_name,
                    "Latitude": round(unit.stop_key[0], 6), "Longitude": round(unit.stop_key[1], 6),
                    "Bus Reaches Stop": minutes_to_clock(reach),
                    "Boarding (min)": round(unit.dwell_min, 1),
                    "Passengers Boarding": unit.pax,
                    "Ride of These Passengers (min)": round(ride, 1),
                    "Drive to Next Stop (min)": round(M[unit.node][seq_nodes[pos + 1] if pos + 1 < len(route) else FAC], 1),
                    "Cumulative Passengers": sum(units[x].pax for x in route[:pos + 1]),
                    "Flag": flag,
                })
                if flag:
                    over.append(unit.stop_name)
                for emp in unit.employees:
                    passenger_rows.append({
                        "Shift": shift, "Bus ID": bus_id, "Stop Seq": pos + 1, "Employee ID": emp["employee_id"],
                        "Employee Name": emp["employee_name"], "Mobile No": emp["mobile"], "Area / Landmark": emp["area"],
                        "Pickup Point": emp["pickup_raw"], "Bus Reaches Stop": minutes_to_clock(reach),
                        "Ride Time (min)": round(ride, 1), "Flag": flag,
                    })
                    if flag:
                        review.append(_review("Ride time over limit", emp, flag, shift))

            summary_rows.append({
                "Bus ID": bus_id, "Shift": shift, "Passengers": pax_total,
                "Seats Free": p.bus_capacity - pax_total, "Stops": len(route),
                "Dispatch from Depot": minutes_to_clock(depart) if DEP is not None else "(starts at first stop)",
                "Bus Reaches First Stop": minutes_to_clock(arrival - route_min),
                "Reaches Facility": minutes_to_clock(arrival),
                "Route Duration (min)": round(route_min, 1),
                "Longest Ride (min)": round(max(rides), 1),
                "Average Ride (min)": round(pax_min / pax_total, 1) if pax_total else 0,
                "Road Distance (km)": round(road_km, 1),
                "First Stop": units[route[0]].stop_name, "Last Stop Before Facility": units[route[-1]].stop_name,
                "Status": "REVIEW: " + "; ".join(over) if over else "OK",
            })

    # Unverified employees are listed for review only; they never enter routing.
    summary = pd.DataFrame(summary_rows)
    stop_df = pd.DataFrame(stop_rows)
    passenger_df = pd.DataFrame(passenger_rows)
    review_df = pd.DataFrame(review)
    if not review_df.empty:
        review_df = review_df.drop_duplicates().reset_index(drop=True)

    _write_workbook(out_path, summary, stop_df, passenger_df, review_df, settings, matrix.source, p)
    _print_console(summary, review_df, out_path, matrix.source)


def _write_workbook(path, summary, stops, passengers, review, settings, source, p):
    path.parent.mkdir(parents=True, exist_ok=True)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    fac = settings["facility"]
    rules = pd.DataFrame([
        ("Facility", f"{fac['name']} ({fac['latitude']}, {fac['longitude']})"),
        ("Bus capacity", p.bus_capacity),
        ("Max ride time (min)", p.max_ride_minutes),
        ("Fixed cost per bus (min equivalent)", p.bus_fixed_cost_minutes),
        ("Boarding time", f"{p.boarding_seconds_per_passenger:.0f} s per passenger, min {p.min_boarding_minutes} min"),
        ("Arrival buffer before shift (min)", settings["arrival_buffer_minutes"]),
        ("Shift start times", ", ".join(f"{k} {v}" for k, v in settings["shifts"].items())),
        ("Travel time source", source),
    ], columns=["Setting", "Value"])

    sheets = [("Rules Used", rules), ("Bus Summary", summary), ("Stop Sequence", stops),
              ("Passenger List", passengers), ("Needs Review", review)]
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets:
            df.to_excel(writer, sheet_name=name, index=False)
        for name, _ in sheets:
            ws = writer.sheets[name]
            for cell in ws[1]:
                cell.font, cell.fill = header_font, header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col_idx in range(1, ws.max_column + 1):
                letter = get_column_letter(col_idx)
                longest = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=10)
                ws.column_dimensions[letter].width = min(max(longest + 2, 10), 60)


def _print_console(summary, review, out_path, source):
    print(f"\nTravel times: {source}")
    if summary.empty:
        print("No buses planned. Check the review list.")
    else:
        print(summary[["Bus ID", "Passengers", "Stops", "Longest Ride (min)", "Status"]].to_string(index=False))
    print("\nNeeds Review (by category):")
    if review.empty:
        print("  none")
    else:
        for category, count in review["Category"].value_counts().items():
            print(f"  {count:4d}  {category}")
    print(f"Written: {out_path}")


def main(argv=None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--settings", type=Path, default=ROOT / "settings.json")
    ap.add_argument("--roster", help="Raw roster Excel file (overrides settings.json)")
    ap.add_argument("--roster-sheet", help="Sheet name for the roster")
    ap.add_argument("--stops", help="Verified pickup-point master Excel file")
    ap.add_argument("--stops-sheet", help="Sheet name for the stop master")
    ap.add_argument("--out", help="Output workbook")
    args = ap.parse_args(argv)
    try:
        plan(args)
    except (ValueError, FileNotFoundError, KeyError) as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
