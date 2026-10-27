"""Shared helpers: text cleaning, pickup-point standardisation, distance, Excel writing."""
import difflib
import math
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

EARTH_RADIUS_KM = 6371.0088


def is_blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return str(value).strip() == ""


def clean_spaces(value) -> str:
    """Trim, collapse runs of whitespace (incl. non-breaking spaces)."""
    if is_blank(value):
        return ""
    text = str(value).replace(" ", " ").replace("\t", " ")
    return re.sub(r"\s+", " ", text).strip()


def smart_case(text: str) -> str:
    """Title-case text that is entirely upper or entirely lower case; leave mixed case alone."""
    if text and (text.isupper() or text.islower()):
        return text.title()
    return text


def normalize_key(text) -> str:
    """Lower-case, letters/digits only, single spaces. Used to match equal names."""
    text = clean_spaces(text).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


# ---------------------------------------------------------------------------
# Reference files
# ---------------------------------------------------------------------------
def _read_csv_if_exists(path: Path) -> pd.DataFrame:
    if path.exists():
        return pd.read_csv(path, dtype=str, keep_default_na=False)
    return pd.DataFrame()


_WORD_CORRECTIONS = None
_ALIASES = None


def word_corrections() -> dict:
    global _WORD_CORRECTIONS
    if _WORD_CORRECTIONS is None:
        df = _read_csv_if_exists(config.WORD_CORRECTIONS_FILE)
        _WORD_CORRECTIONS = {
            clean_spaces(w).lower(): clean_spaces(c).lower()
            for w, c in zip(df.get("wrong", []), df.get("correct", []))
            if clean_spaces(w)
        }
    return _WORD_CORRECTIONS


def stop_aliases() -> dict:
    """Map of normalized stop key -> canonical display name (merges spelling variants)."""
    global _ALIASES
    if _ALIASES is None:
        df = _read_csv_if_exists(config.ALIASES_FILE)
        _ALIASES = {
            normalize_key(a): clean_spaces(c)
            for a, c in zip(df.get("alias", []), df.get("canonical", []))
            if clean_spaces(a)
        }
    return _ALIASES


# ---------------------------------------------------------------------------
# Pickup-point standardisation
# ---------------------------------------------------------------------------
def standardize_pickup(raw) -> str:
    """Return a clean, display-ready pickup point name from a raw bus-stop string."""
    text = clean_spaces(raw)
    if not text:
        return ""
    text = re.sub(r"\s*,\s*", ", ", text)       # "a,b" -> "a, b"
    text = re.sub(r"(,\s*)+", ", ", text)       # ",," -> ","
    text = re.sub(r"\(\s*", "(", text)
    text = re.sub(r"\s*\)", ")", text)
    text = text.strip(" ,.;-")
    text = smart_case(text)

    corrections = word_corrections()

    def fix_word(match):
        word = match.group(0)
        fixed = corrections.get(word.lower())
        if fixed is None:
            return word
        return fixed.title() if word[:1].isupper() else fixed

    text = re.sub(r"[A-Za-z]+", fix_word, text)

    alias = stop_aliases().get(normalize_key(text))
    return alias if alias else text


def pickup_key(display_name: str) -> str:
    return normalize_key(display_name)


def find_similar_stops(keys_to_names: dict, threshold: float) -> list:
    """Return (name_a, name_b, score) for stop names that look like the same place."""
    items = sorted(keys_to_names.items())
    pairs = []
    for i in range(len(items)):
        for j in range(i + 1, len(items)):
            ka, na = items[i]
            kb, nb = items[j]
            if ka == kb:
                continue
            score = difflib.SequenceMatcher(None, ka, kb).ratio()
            if score >= threshold:
                pairs.append((na, nb, round(score, 3)))
    return pairs


# ---------------------------------------------------------------------------
# Distance
# ---------------------------------------------------------------------------
def haversine_km(lat1, lon1, lat2, lon2) -> float:
    """Straight-line (great-circle) distance in km. NOT the road distance."""
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = p2 - p1
    dlmb = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlmb / 2) ** 2
    return 2 * EARTH_RADIUS_KM * math.asin(min(1.0, math.sqrt(a)))


def in_pune_bounds(lat, lon) -> bool:
    b = config.PUNE_BOUNDS
    return b["min_lat"] <= lat <= b["max_lat"] and b["min_lon"] <= lon <= b["max_lon"]


def to_float(value):
    try:
        if is_blank(value):
            return None
        return float(str(value).strip())
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def write_workbook(path: Path, sheets: dict, notes: dict | None = None) -> None:
    """Write several DataFrames to one formatted .xlsx file.

    sheets: {sheet_name: DataFrame}
    notes:  {sheet_name: "one-line description shown in row 1"} (optional)
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    notes = notes or {}
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            safe_name = name[:31]
            start_row = 0
            if notes.get(name):
                pd.DataFrame([[notes[name]]]).to_excel(
                    writer, sheet_name=safe_name, index=False, header=False, startrow=0)
                start_row = 2
            df.to_excel(writer, sheet_name=safe_name, index=False, startrow=start_row)
            ws = writer.sheets[safe_name]
            header_row = start_row + 1
            for cell in ws[header_row]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            if start_row:
                ws.cell(row=1, column=1).font = Font(italic=True, color="555555")
            ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
            if ws.max_row > header_row:
                ws.auto_filter.ref = (
                    f"A{header_row}:{get_column_letter(max(ws.max_column, 1))}{ws.max_row}")
            for idx, col in enumerate(df.columns, start=1):
                values = [str(col)] + [str(v) for v in df[col].head(500).tolist()]
                width = min(max(len(v) for v in values) + 2, 60)
                ws.column_dimensions[get_column_letter(idx)].width = max(width, 8)


def read_sheet(path: Path, sheet: str) -> pd.DataFrame:
    """Read a pipeline sheet (header in row 1, no note rows)."""
    df = pd.read_excel(path, sheet_name=sheet, dtype=object)
    return df.dropna(how="all")
