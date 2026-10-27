"""Tests for verified pickup geocoding and non-blocking roster routing."""
import io
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

import app  # noqa: E402
import auto_resolve  # noqa: E402
import config  # noqa: E402
import geocoding  # noqa: E402


class PickupGeocoding(unittest.TestCase):
    def test_coordinate_master_uses_corrections_and_same_coordinate_aliases(self):
        results = geocoding.resolve_pickup_points(
            ["Xarbia,ambi", "Indrayani collage talegaon", "KATVI ROAD TALEGAON"]
        )
        for result in results.values():
            self.assertEqual(result["location_status"], "VERIFIED")
            self.assertIsNotNone(result["latitude"])
            self.assertIsNotNone(result["longitude"])

    def test_talegaon_low_confidence_candidate_stays_unverified(self):
        query = "Talegaon, Pune district, Maharashtra, India"
        cache = {
            query: {
                "candidates": [
                    {
                        "formatted": "Talegaon Dabhade, MH, India",
                        "lat": 18.7337107,
                        "lon": 73.6728593,
                        "confidence": 0.25,
                        "result_type": "city",
                    }
                ]
            }
        }
        with patch.object(geocoding, "load_cache", return_value=cache), \
             patch.object(geocoding, "_load_coordinate_index", return_value={}):
            result = geocoding.resolve_pickup_points(["Talegaon"])["Talegaon"]
        self.assertEqual(result["location_status"], "LOCATION_UNVERIFIED")
        self.assertIsNone(result["latitude"])

    def test_city_centroid_is_not_used_as_a_pickup_stop(self):
        query = "Talegaon, Pune district, Maharashtra, India"
        cache = {
            query: {
                "candidates": [
                    {
                        "formatted": "Talegaon Dabhade, MH, India",
                        "lat": 18.7337107,
                        "lon": 73.6728593,
                        "confidence": 0.99,
                        "result_type": "city",
                    }
                ]
            }
        }
        with patch.object(geocoding, "load_cache", return_value=cache), \
             patch.object(geocoding, "_load_coordinate_index", return_value={}):
            result = geocoding.resolve_pickup_points(["Talegaon"])["Talegaon"]
        self.assertEqual(result["location_status"], "LOCATION_UNVERIFIED")

    def test_wrong_locality_candidate_is_rejected_even_inside_pune_bounds(self):
        canonical = "Vadgaon, Khandoba Mandir"
        query = canonical + config.GEOAPIFY_QUERY_SUFFIX
        cache = {
            query: {
                "candidates": [
                    {
                        "formatted": "Khandoba Mandir, Mhalunge-Baner Road, Balewadi, Pune",
                        "lat": 18.5692765,
                        "lon": 73.7650538,
                        "confidence": 1,
                    }
                ]
            }
        }
        with patch.object(geocoding, "load_cache", return_value=cache), \
             patch.object(geocoding, "_load_coordinate_index", return_value={}):
            result = geocoding.resolve_pickup_points(
                ["VADGAV, KHANDOBA MANDIR"]
            )["VADGAV, KHANDOBA MANDIR"]
        self.assertEqual(result["location_status"], "LOCATION_UNVERIFIED")
        self.assertIsNone(result["longitude"])

    def test_high_confidence_relevant_search_result_is_cached_and_accepted(self):
        candidate = {
            "formatted": "Talegaon Dabhade, Maharashtra, India",
            "lat": 18.7337107,
            "lon": 73.6728593,
            "confidence": 0.9,
            "result_type": "amenity",
        }
        with (
            patch.object(geocoding, "load_cache", return_value={}),
            patch.object(geocoding, "_load_coordinate_index", return_value={}),
            patch.object(geocoding.config, "GEOAPIFY_API_KEY", "test-key"),
            patch.object(geocoding, "geoapify_lookup", return_value=[candidate]),
            patch.object(geocoding, "save_cache") as save_cache,
        ):
            result = geocoding.resolve_pickup_points(["Talegaon"])["Talegaon"]
        self.assertEqual(result["location_status"], "VERIFIED")
        self.assertEqual(result["coordinate_source"], "Geoapify")
        save_cache.assert_called_once()

    def test_auto_resolver_does_not_approve_unverified_or_fallback_coordinates(self):
        master = pd.DataFrame(
            [
                {
                    "pickup_point": "Old Fallback",
                    "status": "Approved",
                    "latitude": 18.7,
                    "longitude": 73.7,
                    "reviewed_by": "auto_local_fallback",
                },
                {
                    "pickup_point": "Unmatched Stop",
                    "status": "Needs review",
                    "latitude": None,
                    "longitude": None,
                    "reviewed_by": None,
                },
            ]
        )
        unresolved = {
            name: {
                "location_status": "LOCATION_UNVERIFIED",
                "latitude": None,
                "longitude": None,
                "resolution_reason": "No unique verified match.",
            }
            for name in ("Old Fallback", "Unmatched Stop")
        }
        with (
            patch.object(auto_resolve, "read_sheet", return_value=pd.DataFrame()),
            patch.object(auto_resolve, "load_master", return_value=master.copy()),
            patch.object(auto_resolve, "resolve_pickup_points", return_value=unresolved),
            patch.object(auto_resolve, "save_master") as save_master,
        ):
            auto_resolve.resolve_all()

        saved = save_master.call_args.args[0]
        self.assertEqual(saved["status"].tolist(), ["Needs review", "Needs review"])
        self.assertTrue(saved["latitude"].isna().all())
        self.assertTrue(saved["longitude"].isna().all())


class UnverifiedRosterFlow(unittest.TestCase):
    def test_unknown_employee_is_retained_but_not_routed(self):
        frame = pd.DataFrame(
            [
                {
                    "Name": "Routed Employee",
                    "EMP ID": "1",
                    "Shift": "A",
                    "Bus Stop": "Known Stop",
                    "Mob. No.": "111",
                    "GRADE": "G1",
                    "Location": "West",
                },
                {
                    "Name": "Review Employee",
                    "EMP ID": "2",
                    "Shift": "A",
                    "Bus Stop": "Unknown Stop",
                    "Mob. No.": "222",
                    "GRADE": "G2",
                    "Location": "West",
                },
            ]
        )
        locations = {
            "Known Stop": {
                "location_status": "VERIFIED",
                "latitude": 18.735,
                "longitude": 73.675,
            },
            "Unknown Stop": {
                "location_status": "LOCATION_UNVERIFIED",
                "latitude": None,
                "longitude": None,
                "resolution_reason": "No matching verified location.",
            },
        }
        with patch.object(app, "resolve_pickup_points", return_value=locations):
            result = app.process_and_group(frame)
        self.assertEqual(len(result["routes"]), 1)
        self.assertEqual(result["total_employees"], 2)
        self.assertEqual(result["routed_employees"], 1)
        self.assertEqual(result["unverified_employees"], 1)
        unresolved = next(
            employee
            for employee in result["employees"]
            if employee["location_status"] == "LOCATION_UNVERIFIED"
        )
        self.assertEqual(unresolved["employee_name"], "Review Employee")
        self.assertEqual(unresolved["assigned_bus"], "Unassigned")

    def test_resolve_pickups_api_validates_payload(self):
        client = app.app.test_client()
        response = client.post("/api/resolve_pickups", json={"pickup_points": ["Xarbia,ambi"]})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.get_json()["locations"]["Xarbia,ambi"]["location_status"],
            "VERIFIED",
        )
        invalid = client.post("/api/resolve_pickups", json={"pickup_points": [1]})
        self.assertEqual(invalid.status_code, 400)

    def test_excel_export_preserves_unverified_employee_and_reason(self):
        client = app.app.test_client()
        response = client.post(
            "/api/export_excel",
            json={
                "routes": [],
                "employees": [
                    {
                        "sr_no": 1,
                        "employee_id": "2",
                        "employee_name": "Review Employee",
                        "shift": "A",
                        "assigned_bus": "Unassigned",
                        "pickup_point": "Talegaon",
                        "mobile": "222",
                        "location_status": "LOCATION_UNVERIFIED",
                        "resolution_reason": "Ambiguous locality.",
                    }
                ],
            },
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data), read_only=True)
        self.assertIn("Location Review", workbook.sheetnames)
        review = workbook["Location Review"]
        values = list(review.values)
        self.assertTrue(any("LOCATION_UNVERIFIED" in row for row in values))
        self.assertTrue(any("Ambiguous locality." in row for row in values))


if __name__ == "__main__":
    unittest.main()
